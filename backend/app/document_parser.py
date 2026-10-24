from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ALLOWED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf"}


class DocumentParseError(ValueError):
    """Raised when an uploaded document cannot be converted to reviewable text."""


class UnsupportedDocumentTypeError(DocumentParseError):
    """Raised when the uploaded file extension is unsupported."""


def validate_file_extension(filename: str) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        supported = ", ".join(sorted(ALLOWED_EXTENSIONS))
        raise UnsupportedDocumentTypeError(
            f"対応していないファイル形式です。対応形式: {supported}"
        )
    return extension


def extract_document_text(filename: str, content: bytes) -> str:
    extension = validate_file_extension(filename)
    if not content:
        raise DocumentParseError("ファイルが空です。内容のある設計書を選択してください。")

    if extension in {".txt", ".md"}:
        text = _extract_plain_text(content)
    elif extension == ".docx":
        text = _extract_docx(content)
    else:
        text = _extract_pdf(content)

    cleaned_text = text.replace("\x00", "").strip()
    if not cleaned_text:
        if extension == ".pdf":
            raise DocumentParseError(
                "无法从该 PDF 中提取文本，目前暂不支持扫描版 PDF。"
            )
        raise DocumentParseError("文書からレビュー可能なテキストを抽出できませんでした。")
    return cleaned_text


def _extract_plain_text(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentParseError(
            "TXT / Markdown ファイルはUTF-8形式で保存してください。"
        ) from exc


def _extract_docx(content: bytes) -> str:
    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise DocumentParseError("DOCXファイルを読み取れませんでした。") from exc

    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_pdf(content: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise DocumentParseError("PDFファイルを読み取れませんでした。") from exc
