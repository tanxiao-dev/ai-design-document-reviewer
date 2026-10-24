from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from app.document_parser import (
    DocumentParseError,
    UnsupportedDocumentTypeError,
    extract_document_text,
)


def test_extract_txt() -> None:
    assert extract_document_text("design.txt", "基本設計書".encode()) == "基本設計書"


def test_extract_markdown() -> None:
    content = "# 概要\n\n処理仕様です。".encode()
    assert extract_document_text("design.md", content) == content.decode()


def test_extract_docx() -> None:
    document = Document()
    document.add_heading("画面設計", level=1)
    document.add_paragraph("ユーザーIDを入力する。")
    buffer = BytesIO()
    document.save(buffer)

    text = extract_document_text("design.docx", buffer.getvalue())

    assert "画面設計" in text
    assert "ユーザーIDを入力する。" in text


def test_reject_unsupported_file_type() -> None:
    with pytest.raises(UnsupportedDocumentTypeError, match="対応していない"):
        extract_document_text("design.xlsx", b"data")


def test_reject_empty_file() -> None:
    with pytest.raises(DocumentParseError, match="ファイルが空"):
        extract_document_text("design.txt", b"")


def test_reject_pdf_without_extractable_text() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = BytesIO()
    writer.write(buffer)

    with pytest.raises(DocumentParseError, match="扫描版 PDF"):
        extract_document_text("scan.pdf", buffer.getvalue())
